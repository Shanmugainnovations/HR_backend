import os
from employees.permissions import HasRoleAndDataPermission
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework import status
from employees.models import LeaveRequest, Employee, Shift, EmployeeShiftSchedule, LeaveType, Register
from employees.serializers import LeaveTypeSerializer
from employees.views.authentication.auth import resolve_department_names
from django.utils import timezone
from datetime import datetime, timedelta

def resolve_leave_department(leave_obj, reg_map=None):
    raw_dept = getattr(leave_obj, 'department', '') or getattr(leave_obj, 'department_id', '')
    if (not raw_dept or str(raw_dept).strip() == '') and reg_map:
        raw_dept = reg_map.get(str(getattr(leave_obj, 'employee_id', '')).strip(), '')
    
    if not raw_dept:
        return "General"
    
    resolved = resolve_department_names(str(raw_dept))
    if not resolved or resolved == "Unassigned":
        return str(raw_dept) if not str(raw_dept).upper().startswith("DEPT") else "General"
    return resolved

@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
def leave_type_list_create(request):

    if request.method == 'GET':
        leave_types = LeaveType.objects.all().order_by('name')
        serializer = LeaveTypeSerializer(leave_types, many=True)
        return Response(serializer.data)

    elif request.method == 'POST':
        serializer = LeaveTypeSerializer(data=request.data)
        if serializer.is_valid():
            from employees.models import extract_actor_id
            actor_id = extract_actor_id(request)
            lt = serializer.save(created_by=actor_id, lastmodified_by=actor_id)
            print(f"🔑 AUDIT SAVED -> LeaveType CreatedBy: {lt.created_by}")
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([AllowAny])
def leave_type_detail(request, pk):
    try:
        try:
            leave_type = LeaveType.objects.get(pk=int(pk))
        except (ValueError, TypeError):
            leave_type = LeaveType.objects.get(pk=pk)
    except LeaveType.DoesNotExist:
        return Response({'error': 'Leave type not found'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'GET':
        serializer = LeaveTypeSerializer(leave_type)
        return Response(serializer.data)

    elif request.method == 'PUT':
        serializer = LeaveTypeSerializer(leave_type, data=request.data, partial=True)
        if serializer.is_valid():
            lt = serializer.save()
            lt.save_with_audit(request)
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == 'DELETE':
        leave_type.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

@api_view(['POST'])
@permission_classes([AllowAny])
def apply_leave(request):
    try:
        employee_id = getattr(request, 'authenticated_employee_id', None) or request.data.get('employee_id')
        if not employee_id:
            return Response({"error": "Employee ID is required"}, status=400)
            
        reg = Register.objects.filter(employee_id=str(employee_id).strip()).first()
        emp_name = request.data.get('employee_name') or (reg.name if reg else None)
        raw_dept = request.data.get('department') or (reg.department if reg else None)
        raw_dept_id = request.data.get('department_id') or (reg.department if reg else None)

        # Fallback to MongoDB profile if reg is missing or department is missing
        if not reg or not raw_dept:
            try:
                from employees.views.common.utils import get_mongo_client
                client = get_mongo_client()
                if client:
                    global_db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]
                    prof = global_db['backend_diagnostics_profile'].find_one(
                        {'$or': [{'employeeId': str(employee_id).strip()}, {'employee_id': str(employee_id).strip()}]}
                    )
                    if prof:
                        if not emp_name:
                            fn = prof.get('firstName') or ''
                            ln = prof.get('lastName') or ''
                            emp_name = f"{fn} {ln}".strip() or prof.get('name')
                        if not raw_dept:
                            raw_dept = prof.get('department')
                        if not raw_dept_id:
                            raw_dept_id = prof.get('department')
            except Exception:
                pass

        dept_name = resolve_department_names(raw_dept) if raw_dept else ""
        if not dept_name or dept_name == "Unassigned":
            dept_name = resolve_department_names(raw_dept_id) if raw_dept_id else ""
        if not dept_name or dept_name == "Unassigned":
            dept_name = raw_dept or "General"

        # Resolve requester role robustly from request headers, token, register, or payload
        from employees.views.common.utils import get_request_user_approval_context, resolve_department_filter
        ctx = get_request_user_approval_context(request)
        requester_role = str(
            request.data.get('role') or
            request.headers.get('User-Role') or
            request.headers.get('X-User-Role') or
            ctx.get('role_str') or
            (reg.role if reg else '')
        ).strip().upper()

        # 3-Tier Hierarchical Approval routing:
        if 'HOD' in requester_role:
            # 1. HOD applies -> Routes directly to Top Management (AVP)!
            initial_status = 'Pending AVP'
        elif 'INCHARGE' in requester_role or 'HR-R-INC' in requester_role or requester_role.endswith('-INC'):
            # 2. Incharge applies -> Routes directly to HOD!
            initial_status = 'Pending HOD'
        else:
            # 3. Regular employee: check if Incharge exists for this department
            has_incharge = False
            dept_key = raw_dept_id or raw_dept
            if dept_key:
                dept_ctx = resolve_department_filter(dept_key)
                target_terms = dept_ctx.get('target_terms', set()) or {dept_key}
                from django.db.models import Q
                q_dept = Q()
                for t in target_terms:
                    q_dept |= Q(department__icontains=t) | Q(assigned_departments__icontains=t)
                inc_user = Register.objects.filter(
                    Q(role__in=['HR-R-INC', 'Incharge', 'INCHARGE']) | Q(role__icontains='INCHARGE') | Q(role__icontains='HR-R-INC'),
                    q_dept
                ).first()
                if inc_user:
                    has_incharge = True
            # If department has an Incharge -> Level 1 (Pending Incharge)
            # If no Incharge allocated for this department -> Routes directly to Level 2 (Pending HOD)
            initial_status = 'Pending Incharge' if has_incharge else 'Pending HOD'

        LeaveRequest.objects.create(
            employee_id=employee_id,
            employee_name=emp_name,
            department=dept_name,
            department_id=raw_dept_id or raw_dept or dept_name,
            start_date=request.data.get('start_date'),
            end_date=request.data.get('end_date'),
            leave_type=request.data.get('leave_type'),
            reason=request.data.get('reason'),
            status=initial_status
        )
        return Response({
            "message": "Leave requested successfully",
            "status": initial_status
        }, status=201)
    except Exception as e:
        return Response({"error": str(e)}, status=500)

@api_view(['GET'])
@permission_classes([AllowAny])
def my_leaves(request):
    try:
        employee_id = getattr(request, 'authenticated_employee_id', None) or request.GET.get('employee_id')
        if not employee_id:
            return Response({"error": "Employee ID is required"}, status=400)
            
        leaves = LeaveRequest.objects.filter(employee_id=employee_id).order_by('-applied_on')
        reg = Register.objects.filter(employee_id=str(employee_id).strip()).first()
        reg_map = {str(employee_id).strip(): reg.department} if reg and reg.department else {}

        data = [{
            "id": l.id,
            "employee_id": l.employee_id,
            "employee_name": l.employee_name or (reg.name if reg else "Employee"),
            "department": resolve_leave_department(l, reg_map),
            "department_code": l.department_id or l.department,
            "start_date": l.start_date,
            "end_date": l.end_date,
            "leave_type": l.leave_type,
            "reason": l.reason,
            "status": l.status,
            "applied_on": l.applied_on,
            "reviewed_by_name": l.reviewed_by_name,
            "incharge_name": getattr(l, 'incharge_name', None),
            "hod_name": getattr(l, 'hod_name', None),
            "avp_name": getattr(l, 'avp_name', None),
        } for l in leaves]
        return Response(data, status=200)
    except Exception as e:
        return Response({"error": str(e)}, status=500)

@api_view(['GET'])
@permission_classes([AllowAny])
def pending_leaves(request):
    try:
        from django.db.models import Q
        from employees.views.common.utils import get_request_user_approval_context, resolve_department_filter

        department_id = request.GET.get('department_id')
        department = request.GET.get('department')
        status_filter = request.GET.get('status')
        
        leaves = LeaveRequest.objects.all().order_by('-applied_on')

        ctx = get_request_user_approval_context(request)
        is_avp = ctx['is_avp']
        is_hod = ctx['is_hod']
        is_incharge = ctx['is_incharge']
        is_admin = ctx['is_admin']
        caller_id = ctx['employee_id']
        assigned_depts = ctx['assigned_departments']

        if is_avp:
            # 1. AVP Level: View HOD leaves (Pending AVP)
            if status_filter and status_filter != 'All':
                leaves = leaves.filter(status__iexact=status_filter)
            else:
                leaves = leaves.filter(status__in=['Pending AVP', 'Pending HOD', 'Pending'])
            if department and department != 'All':
                dept_ctx = resolve_department_filter(department)
                target_terms = dept_ctx['target_terms']
                q_dept = Q()
                for t in target_terms:
                    q_dept |= Q(department__icontains=t) | Q(department_id__icontains=t)
                leaves = leaves.filter(q_dept)

        elif is_hod:
            # 2. HOD Level: View department leaves waiting for HOD approval
            if caller_id:
                leaves = leaves.exclude(employee_id=str(caller_id).strip())

            if status_filter and status_filter != 'All':
                leaves = leaves.filter(status__iexact=status_filter)
            else:
                leaves = leaves.filter(status__in=['Pending HOD', 'Pending'])

            target_depts = assigned_depts or ([department] if department and department != 'All' else [])
            if target_depts:
                q_hod = Q()
                for d in target_depts:
                    dept_ctx = resolve_department_filter(d)
                    target_terms = dept_ctx['target_terms']
                    matching_emp_ids = dept_ctx['matching_employee_ids'] or set()
                    if matching_emp_ids:
                        q_hod |= Q(employee_id__in=matching_emp_ids)
                    for t in target_terms:
                        q_hod |= Q(department__icontains=t) | Q(department_id__icontains=t)
                leaves = leaves.filter(q_hod)
            else:
                leaves = LeaveRequest.objects.none()

        elif is_incharge:
            # 3. INCHARGE Level: View ward/department staff leaves for Level 1 review
            if caller_id:
                leaves = leaves.exclude(employee_id=str(caller_id).strip())

            if status_filter and status_filter != 'All':
                leaves = leaves.filter(status__iexact=status_filter)
            else:
                leaves = leaves.filter(status__in=['Pending Incharge', 'Pending'])

            target_depts = assigned_depts or ([department] if department and department != 'All' else [])
            if target_depts:
                q_inc = Q()
                for d in target_depts:
                    dept_ctx = resolve_department_filter(d)
                    target_terms = dept_ctx['target_terms']
                    matching_emp_ids = dept_ctx['matching_employee_ids'] or set()
                    if matching_emp_ids:
                        q_inc |= Q(employee_id__in=matching_emp_ids)
                    for t in target_terms:
                        q_inc |= Q(department__icontains=t) | Q(department_id__icontains=t)
                leaves = leaves.filter(q_inc)
            else:
                leaves = LeaveRequest.objects.none()

        elif is_admin:
            # 4. ADMIN: Full visibility
            if status_filter and status_filter != 'All':
                leaves = leaves.filter(status__iexact=status_filter)
            elif not status_filter:
                leaves = leaves.filter(status__in=['Pending', 'Pending Incharge', 'Pending HOD', 'Pending AVP'])
            if department and department != 'All':
                dept_ctx = resolve_department_filter(department)
                target_terms = dept_ctx['target_terms']
                q_dept = Q()
                for t in target_terms:
                    q_dept |= Q(department__icontains=t) | Q(department_id__icontains=t)
                leaves = leaves.filter(q_dept)
        else:
            # Non-approver employee has no approval items
            return Response([], status=200)

        emp_ids = [str(l.employee_id).strip() for l in leaves if l.employee_id]
        reg_map = {
            str(r.employee_id).strip(): (r.department or '')
            for r in Register.objects.filter(employee_id__in=emp_ids)
        }

        data = [{
            "id": l.id,
            "employee_id": l.employee_id,
            "employee_name": l.employee_name or f"Employee #{l.employee_id}",
            "department": resolve_leave_department(l, reg_map),
            "department_code": l.department_id or l.department,
            "start_date": l.start_date,
            "end_date": l.end_date,
            "leave_type": l.leave_type,
            "reason": l.reason,
            "status": l.status,
            "applied_on": l.applied_on,
            "reviewed_by_name": l.reviewed_by_name,
            # Hierarchical approval fields
            "incharge_id": getattr(l, 'incharge_id', None),
            "incharge_name": getattr(l, 'incharge_name', None),
            "incharge_action": getattr(l, 'incharge_action', None),
            "incharge_remarks": getattr(l, 'incharge_remarks', None),
            "incharge_date": getattr(l, 'incharge_date', None),
            "hod_id": getattr(l, 'hod_id', None),
            "hod_name": getattr(l, 'hod_name', None),
            "hod_action": getattr(l, 'hod_action', None),
            "hod_remarks": getattr(l, 'hod_remarks', None),
            "hod_date": getattr(l, 'hod_date', None),
            "avp_id": getattr(l, 'avp_id', None),
            "avp_name": getattr(l, 'avp_name', None),
            "avp_action": getattr(l, 'avp_action', None),
            "avp_remarks": getattr(l, 'avp_remarks', None),
            "avp_date": getattr(l, 'avp_date', None),
        } for l in leaves]
        return Response(data, status=200)
    except Exception as e:
        return Response({"error": str(e)}, status=500)

@api_view(['GET'])
@permission_classes([AllowAny])
def leave_history(request):
    try:
        status_filter = request.GET.get('status')
        employee_name = request.GET.get('employee_name')
        department = request.GET.get('department')
        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')
        
        leaves = LeaveRequest.objects.all()
        
        if status_filter and status_filter != 'All':
            leaves = leaves.filter(status=status_filter)
        if employee_name:
            leaves = leaves.filter(employee_name__icontains=employee_name)
        if department and department != 'All':
            from django.db.models import Q
            from employees.views.common.utils import resolve_department_filter
            dept_ctx = resolve_department_filter(department)
            target_terms = dept_ctx['target_terms']
            matching_emp_ids = dept_ctx['matching_employee_ids'] or set()
            
            q_dept = Q()
            if matching_emp_ids:
                q_dept |= Q(employee_id__in=matching_emp_ids)
            for t in target_terms:
                q_dept |= Q(department__icontains=t) | Q(department_id__icontains=t)
            leaves = leaves.filter(q_dept)
        if from_date:
            leaves = leaves.filter(end_date__gte=from_date)
        if to_date:
            leaves = leaves.filter(start_date__lte=to_date)
            
        leaves = leaves.order_by('-applied_on')
        
        hist_emp_ids = [str(l.employee_id).strip() for l in leaves if l.employee_id]
        hist_reg_map = {
            str(r.employee_id).strip(): (r.department or '')
            for r in Register.objects.filter(employee_id__in=hist_emp_ids)
        }

        data = [{
            "id": l.id,
            "employee_id": l.employee_id,
            "employee_name": l.employee_name or f"Employee #{l.employee_id}",
            "department": resolve_leave_department(l, hist_reg_map),
            "department_code": l.department_id or l.department,
            "start_date": l.start_date,
            "end_date": l.end_date,
            "leave_type": l.leave_type,
            "reason": l.reason,
            "status": l.status,
            "applied_on": l.applied_on,
            "reviewed_by_name": l.reviewed_by_name,
            "incharge_name": getattr(l, 'incharge_name', None),
            "incharge_remarks": getattr(l, 'incharge_remarks', None),
            "hod_name": getattr(l, 'hod_name', None),
            "hod_remarks": getattr(l, 'hod_remarks', None),
            "avp_name": getattr(l, 'avp_name', None),
            "avp_remarks": getattr(l, 'avp_remarks', None),
        } for l in leaves]
        return Response(data, status=200)
    except Exception as e:
        return Response({"error": str(e)}, status=500)

@api_view(['PUT', 'POST'])
@permission_classes([AllowAny])
def update_leave_status(request, leave_id):
    try:
        from employees.views.common.utils import get_request_user_approval_context
        from django.utils import timezone

        status_val = request.data.get('status')
        reviewer_name = request.data.get('reviewer_name') or 'Reviewer'
        remarks = request.data.get('remarks') or request.data.get('admin_remarks') or ''
        reviewer_role = str(request.data.get('reviewer_role') or '').strip().upper()

        if status_val not in ['Approved', 'Rejected']:
            return Response({"error": "Invalid status. Must be 'Approved' or 'Rejected'"}, status=400)

        leave = LeaveRequest.objects.get(id=leave_id)
        prev_status = leave.status
        now_dt = timezone.now()

        ctx = get_request_user_approval_context(request)
        is_incharge = ctx['is_incharge'] or ('INCHARGE' in reviewer_role) or ('HR-R-INC' in reviewer_role) or reviewer_role.endswith('-INC')
        is_hod = (ctx['is_hod'] or ('HOD' in reviewer_role) or ('HR-R-HOD' in reviewer_role)) and not ctx['is_avp']
        is_avp = ctx['is_avp'] or ('AVP' in reviewer_role) or ('HR-R-AVP' in reviewer_role)
        is_admin = ctx['is_admin'] or ('ADMIN' in reviewer_role) or (reviewer_role in ['HR', 'HR-ADMIN'])

        final_approved = False
        action_msg = ""

        if is_incharge and not (is_hod or is_avp or is_admin):
            # Incharge (Level 1 Review):
            leave.incharge_id = ctx['employee_id'] or ''
            leave.incharge_name = reviewer_name
            leave.incharge_action = status_val
            leave.incharge_remarks = remarks
            leave.incharge_date = now_dt
            if status_val == 'Approved':
                leave.status = 'Pending HOD'
                action_msg = f"Leave recommended by Incharge {reviewer_name} and forwarded to HOD."
            else:
                leave.status = 'Rejected'
                action_msg = f"Leave rejected by Incharge {reviewer_name}."

        elif is_hod and not is_avp:
            # HOD (Level 2 or Direct Approver):
            leave.hod_id = ctx['employee_id'] or ''
            leave.hod_name = reviewer_name
            leave.hod_action = status_val
            leave.hod_remarks = remarks
            leave.hod_date = now_dt
            if status_val == 'Approved':
                leave.status = 'Approved'
                final_approved = True
                action_msg = f"Leave approved by HOD {reviewer_name}."
            else:
                leave.status = 'Rejected'
                action_msg = f"Leave rejected by HOD {reviewer_name}."

        elif is_avp:
            # AVP (Top Management Approver for HOD leaves):
            leave.avp_id = ctx['employee_id'] or ''
            leave.avp_name = reviewer_name
            leave.avp_action = status_val
            leave.avp_remarks = remarks
            leave.avp_date = now_dt
            if status_val == 'Approved':
                leave.status = 'Approved'
                final_approved = True
                action_msg = f"HOD leave approved by AVP {reviewer_name}."
            else:
                leave.status = 'Rejected'
                action_msg = f"HOD leave rejected by AVP {reviewer_name}."

        else:
            # Admin or generic fallback
            leave.status = status_val
            if status_val == 'Approved':
                final_approved = True
            action_msg = f"Leave {status_val.lower()} by {reviewer_name}."

        leave.reviewed_by_name = reviewer_name
        leave.save()
        
        # Update roster ONLY if final approved
        if final_approved:
            try:
                start_date_obj = leave.start_date
                end_date_obj = leave.end_date
                
                leave_shift = Shift.objects.filter(name__iexact=leave.leave_type).first()
                employee_obj = Employee.objects.filter(employee_id=leave.employee_id).first()
                
                if leave_shift and employee_obj:
                    current_date = start_date_obj
                    while current_date <= end_date_obj:
                        sched = EmployeeShiftSchedule.objects.filter(employee=employee_obj, date=current_date).first()
                        if sched:
                            sched.shift = leave_shift
                            sched.save()
                        else:
                            last_sched = EmployeeShiftSchedule.objects.order_by('-id').first()
                            new_id = (last_sched.id + 1) if last_sched else 1
                            EmployeeShiftSchedule.objects.create(
                                id=new_id,
                                employee=employee_obj,
                                date=current_date,
                                shift=leave_shift
                            )
                        current_date += timedelta(days=1)
            except Exception as roster_e:
                import traceback
                traceback.print_exc()
                print("Failed to update roster on leave approval:", roster_e)

        # Synchronize leave balance in Mongo ONLY if final approved
        try:
            if final_approved and prev_status != 'Approved':
                from pymongo import MongoClient
                mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
                client = MongoClient(mongo_uri)
                global_db = client['Global']
                bal_coll = global_db['employees_leave_balance']
                
                emp_id_str = str(leave.employee_id).strip()
                bal_doc = bal_coll.find_one({'employee_id': emp_id_str})
                
                # Calculate days requested
                if leave.start_date and leave.end_date:
                    leave_days = (leave.end_date - leave.start_date).days + 1
                else:
                    leave_days = 1.0

                ltype_upper = str(leave.leave_type or '').upper().strip()
                type_key = 'EL'
                if 'NH' in ltype_upper or 'HOLIDAY' in ltype_upper or 'PH' in ltype_upper:
                    type_key = 'NH'
                elif 'SL' in ltype_upper or 'SICK' in ltype_upper or 'CL' in ltype_upper or 'CASUAL' in ltype_upper or 'C' in ltype_upper:
                    type_key = 'SL'
                elif 'EL' in ltype_upper or 'EARNED' in ltype_upper:
                    type_key = 'EL'

                if bal_doc:
                    lb = bal_doc.get('leave_balances', {})
                    current_val = float(lb.get(type_key, lb.get(type_key.lower(), 0.0)))
                    new_val = max(0.0, current_val - leave_days)
                    lb[type_key] = new_val
                    tot_avail = sum(float(v) for k, v in lb.items() if k != 'total_available')
                    lb['total_available'] = round(tot_avail, 2)
                    bal_coll.update_one({'_id': bal_doc['_id']}, {'$set': {'leave_balances': lb, 'updated_at': datetime.now()}})
        except Exception as bal_err:
            print("Error updating leave balance in Mongo:", bal_err)

        # Auto-create real-time notification for the employee
        try:
            from employees.views.mobile_app.notifications import get_notifications_collection
            col = get_notifications_collection()
            now_dt = datetime.now()
            start_str = leave.start_date.strftime('%d %b %Y') if hasattr(leave.start_date, 'strftime') else str(leave.start_date)
            end_str = leave.end_date.strftime('%d %b %Y') if hasattr(leave.end_date, 'strftime') else str(leave.end_date)
            remark_note = f" Remarks: {remarks}" if remarks else ""

            if is_incharge and status_val == 'Approved':
                title = "Leave Recommended by Incharge ⏱️"
                message = f"Your request for {leave.leave_type} ({start_str} to {end_str}) was recommended by Incharge {reviewer_name} and forwarded to HOD."
            elif final_approved:
                title = "Leave Approved ✅"
                message = f"Your request for {leave.leave_type} ({start_str} to {end_str}) has been approved by {reviewer_name}."
            else:
                title = "Leave Rejected ❌"
                message = f"Your request for {leave.leave_type} ({start_str} to {end_str}) has been rejected by {reviewer_name}.{remark_note}"

            col.insert_one({
                "employee_id": str(leave.employee_id),
                "title": title,
                "message": message,
                "category": "leave",
                "is_read": False,
                "action_url": "",
                "created_at": now_dt.strftime('%Y-%m-%d %H:%M:%S'),
                "created_at_ts": now_dt.timestamp()
            })
        except Exception as notif_err:
            print("Error pushing leave notification", notif_err)
            
        return Response({"message": action_msg or f"Leave {status_val}", "status": leave.status}, status=200)
    except LeaveRequest.DoesNotExist:
        return Response({"error": "Leave request not found"}, status=404)
    except Exception as e:
        return Response({"error": str(e)}, status=500)


@api_view(['GET'])
@permission_classes([AllowAny])
def get_leave_balances(request):
    """
    Retrieve employee leave balances from employees_leave_balance collection.
    Supports query params:
      - employee_id: filter by specific employee ID
      - department: filter by department name or code (or 'All')
      - search: filter by name or ID
      - balance_filter: 'all' | 'has_balance' | 'zero_balance' | 'high_balance'
      - export: 'csv'
    """
    import os
    import csv
    import re
    from pymongo import MongoClient
    from django.http import HttpResponse

    try:
        mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
        client = MongoClient(mongo_uri)
        db = client['Global']
        coll = db['employees_leave_balance']

        emp_id = request.query_params.get('employee_id')
        dept_param = request.query_params.get('department')
        search_query = request.query_params.get('search', '').strip()
        balance_filter = request.query_params.get('balance_filter', 'all')
        export_csv = request.query_params.get('export') == 'csv'

        from employees.views.common.utils import get_request_user_hod_departments
        is_hod, hod_assigned_depts = get_request_user_hod_departments(request)

        # Build mongo query
        query = {}
        if emp_id:
            query['employee_id'] = str(emp_id).strip()

        if is_hod:
            if hod_assigned_depts:
                if dept_param and dept_param.lower() != 'all':
                    dept_escaped = re.escape(dept_param.strip())
                    query['$or'] = [
                        {'department': {'$regex': f'^{dept_escaped}$', '$options': 'i'}},
                        {'department_code': {'$regex': f'^{dept_escaped}$', '$options': 'i'}},
                        {'raw_department': {'$regex': f'^{dept_escaped}$', '$options': 'i'}}
                    ]
                else:
                    dept_or_conds = []
                    for ad in hod_assigned_depts:
                        ad_escaped = re.escape(ad.strip())
                        dept_or_conds.extend([
                            {'department': {'$regex': f'^{ad_escaped}$', '$options': 'i'}},
                            {'department_code': {'$regex': f'^{ad_escaped}$', '$options': 'i'}},
                            {'raw_department': {'$regex': f'^{ad_escaped}$', '$options': 'i'}}
                        ])
                    query['$or'] = dept_or_conds
            else:
                query['employee_id'] = '__NONE__'
        elif dept_param and dept_param.lower() != 'all':
            dept_escaped = re.escape(dept_param.strip())
            query['$or'] = [
                {'department': {'$regex': f'^{dept_escaped}$', '$options': 'i'}},
                {'department_code': {'$regex': f'^{dept_escaped}$', '$options': 'i'}},
                {'raw_department': {'$regex': f'^{dept_escaped}$', '$options': 'i'}}
            ]

        if search_query:
            search_escaped = re.escape(search_query)
            search_cond = [
                {'employee_id': {'$regex': search_escaped, '$options': 'i'}},
                {'employee_name': {'$regex': search_escaped, '$options': 'i'}},
                {'designation': {'$regex': search_escaped, '$options': 'i'}}
            ]
            if '$or' in query:
                query['$and'] = [
                    {'$or': query.pop('$or')},
                    {'$or': search_cond}
                ]
            else:
                query['$or'] = search_cond

        if balance_filter == 'has_balance':
            query['leave_balances.total_available'] = {'$gt': 0}
        elif balance_filter == 'zero_balance':
            query['leave_balances.total_available'] = {'$lte': 0}
        elif balance_filter == 'high_balance':
            query['leave_balances.total_available'] = {'$gte': 10}

        include_inactive = str(request.query_params.get('include_inactive', 'false')).lower() in ['true', '1', 'yes']
        from employees.views.common.utils import get_inactive_employee_ids
        inactive_ids = get_inactive_employee_ids()
        if not include_inactive and inactive_ids:
            if 'employee_id' in query:
                if isinstance(query['employee_id'], str) and query['employee_id'] in inactive_ids:
                    return Response({'summary': {'total_employees': 0, 'total_available_leaves': 0, 'total_el': 0, 'total_nh': 0, 'total_sl': 0, 'filtered_count': 0, 'departments': []}, 'records': []}, status=200)
            else:
                query['employee_id'] = {'$nin': list(inactive_ids)}

        # Fetch records
        cursor = coll.find(query).sort([
            ('leave_balances.total_available', -1),
            ('employee_id', 1)
        ])
        raw_records = list(cursor)

        # Fallback: If querying a specific employee and no balance record exists yet, create default entry from Mongo profile
        if not raw_records and emp_id:
            try:
                prof_doc = db['backend_diagnostics_profile'].find_one({'employeeId': str(emp_id).strip()})
                if prof_doc:
                    default_doc = {
                        'employee_id': str(emp_id).strip(),
                        'employee_name': prof_doc.get('employeeName') or f"{prof_doc.get('first_name', '')} {prof_doc.get('last_name', '')}".strip() or str(emp_id),
                        'department': prof_doc.get('department_name') or prof_doc.get('department') or 'General',
                        'department_code': prof_doc.get('department') or 'DEPT001',
                        'designation': prof_doc.get('designation') or 'Staff',
                        'doj': prof_doc.get('doj') or '2024-01-01',
                        'leave_balances': {'EL': 12.0, 'NH': 9.0, 'SL': 12.0, 'total_available': 33.0},
                        'eligible_entitlement': {'EL': 12.0, 'NH': 9.0, 'SL': 12.0, 'total_eligible': 33.0},
                        'created_at': datetime.now(),
                        'updated_at': datetime.now()
                    }
                    coll.insert_one(default_doc)
                    raw_records = [default_doc]
            except Exception as fb_err:
                print("Error in fallback leave balance creation:", fb_err)

        # Format records and calculate approved leave taken
        records = []
        rec_emp_ids = [str(r.get('employee_id')).strip() for r in raw_records if r.get('employee_id')]
        taken_summary = {}
        if rec_emp_ids:
            try:
                approved_reqs = LeaveRequest.objects.filter(employee_id__in=rec_emp_ids, status='Approved')
                for req in approved_reqs:
                    eid = str(req.employee_id).strip()
                    days = (req.end_date - req.start_date).days + 1 if req.start_date and req.end_date else 1
                    ltype = str(req.leave_type or '').upper().strip()
                    
                    tkey = 'EL'
                    if 'NH' in ltype or 'HOLIDAY' in ltype or 'PH' in ltype:
                        tkey = 'NH'
                    elif 'SL' in ltype or 'SICK' in ltype or 'CL' in ltype or 'CASUAL' in ltype or 'C' in ltype:
                        tkey = 'SL'

                    if eid not in taken_summary:
                        taken_summary[eid] = {'total_taken': 0.0, 'EL': 0.0, 'NH': 0.0, 'SL': 0.0}
                    taken_summary[eid]['total_taken'] += days
                    taken_summary[eid][tkey] = taken_summary[eid].get(tkey, 0.0) + days
            except Exception as t_err:
                print("Error computing taken leaves summary:", t_err)

        for doc in raw_records:
            doc['_id'] = str(doc['_id'])
            if 'created_at' in doc and hasattr(doc['created_at'], 'isoformat'):
                doc['created_at'] = doc['created_at'].isoformat()
            if 'updated_at' in doc and hasattr(doc['updated_at'], 'isoformat'):
                doc['updated_at'] = doc['updated_at'].isoformat()
            
            eid = str(doc.get('employee_id')).strip()
            ee = doc.get('eligible_entitlement', {})
            lb = doc.get('leave_balances', {})

            el_elig = float(ee.get('EL', ee.get('el', 12.0)))
            nh_elig = float(ee.get('NH', ee.get('nh', 9.0)))
            sl_elig = float(ee.get('SL', ee.get('sl', 12.0)))
            tot_elig = float(ee.get('total_eligible', el_elig + nh_elig + sl_elig))

            # Determine taken leaves from sheet (prioritizing explicit leaves_taken document)
            lt_doc = doc.get('leaves_taken', {})
            if lt_doc and ('EL' in lt_doc or 'total_taken' in lt_doc):
                el_sheet_taken = float(lt_doc.get('EL', 0.0))
                nh_sheet_taken = float(lt_doc.get('NH', 0.0))
                sl_sheet_taken = float(lt_doc.get('SL', 0.0))
            else:
                el_sheet_taken = float(lb.get('EL', lb.get('el', 0.0)))
                nh_sheet_taken = float(lb.get('NH', lb.get('nh', 0.0)))
                sl_sheet_taken = float(lb.get('SL', lb.get('sl', 0.0)))

            req_taken = taken_summary.get(eid, {'total_taken': 0.0, 'EL': 0.0, 'NH': 0.0, 'SL': 0.0})

            el_taken_tot = round(el_sheet_taken + req_taken.get('EL', 0.0), 2)
            nh_taken_tot = round(nh_sheet_taken + req_taken.get('NH', 0.0), 2)
            sl_taken_tot = round(sl_sheet_taken + req_taken.get('SL', 0.0), 2)
            total_taken_tot = round(el_taken_tot + nh_taken_tot + sl_taken_tot, 2)

            doc['leaves_taken'] = {
                'total_taken': total_taken_tot,
                'EL': el_taken_tot,
                'NH': nh_taken_tot,
                'SL': sl_taken_tot
            }

            el_rem = max(0.0, round(el_elig - el_taken_tot, 2))
            nh_rem = max(0.0, round(nh_elig - nh_taken_tot, 2))
            sl_rem = max(0.0, round(sl_elig - sl_taken_tot, 2))
            tot_rem = round(el_rem + nh_rem + sl_rem, 2)

            doc['leave_balances'] = {
                'EL': el_rem,
                'NH': nh_rem,
                'SL': sl_rem,
                'total_available': tot_rem
            }
            records.append(doc)

        # Handle CSV export
        if export_csv:
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = f'attachment; filename="Employee_Leave_Balances_{datetime.now().strftime("%Y%m%d")}.csv"'
            writer = csv.writer(response)
            writer.writerow([
                'S.No', 'Employee ID', 'Employee Name', 'Department', 'Department Code',
                'Designation', 'DOJ', 'Total Allocated Leaves', 'Total Leaves Used', 'Total Pending Balance',
                'EL Pending', 'EL Used', 'EL Total',
                'NH Pending', 'NH Used', 'NH Total',
                'SL Pending', 'SL Used', 'SL Total'
            ])
            for idx, r in enumerate(records, 1):
                lb = r.get('leave_balances', {})
                ee = r.get('eligible_entitlement', {})
                lt = r.get('leaves_taken', {})
                writer.writerow([
                    idx,
                    r.get('employee_id', ''),
                    r.get('employee_name', ''),
                    r.get('department', ''),
                    r.get('department_code', ''),
                    r.get('designation', ''),
                    r.get('doj', ''),
                    ee.get('total_eligible', 33),
                    lt.get('total_taken', 0),
                    lb.get('total_available', 0),
                    lb.get('EL', 0),
                    lt.get('EL', 0),
                    ee.get('EL', 12),
                    lb.get('NH', 0),
                    lt.get('NH', 0),
                    ee.get('NH', 9),
                    lb.get('SL', 0),
                    lt.get('SL', 0),
                    ee.get('SL', 12)
                ])
            return response

        # Aggregate summary statistics across active documents in collection
        pipeline = []
        if not include_inactive and inactive_ids:
            pipeline.append({'$match': {'employee_id': {'$nin': list(inactive_ids)}}})
        if is_hod:
            if hod_assigned_depts:
                hod_match_conds = []
                for ad in hod_assigned_depts:
                    ad_escaped = re.escape(ad.strip())
                    hod_match_conds.extend([
                        {'department': {'$regex': f'^{ad_escaped}$', '$options': 'i'}},
                        {'department_code': {'$regex': f'^{ad_escaped}$', '$options': 'i'}},
                        {'raw_department': {'$regex': f'^{ad_escaped}$', '$options': 'i'}}
                    ])
                pipeline.append({'$match': {'$or': hod_match_conds}})
            else:
                pipeline.append({'$match': {'employee_id': '__NONE__'}})

        pipeline.extend([
            {'$group': {
                '_id': '$department',
                'department_code': {'$first': '$department_code'},
                'count': {'$sum': 1},
                'total_leaves': {'$sum': '$leave_balances.total_available'},
                'total_el': {'$sum': '$leave_balances.EL'},
                'total_nh': {'$sum': '$leave_balances.NH'},
                'total_sl': {'$sum': '$leave_balances.SL'},
                'total_eligible': {'$sum': '$eligible_entitlement.total_eligible'},
                'total_taken': {'$sum': '$leaves_taken.total_taken'}
            }},
            {'$sort': {'total_leaves': -1}}
        ])
        dept_summary_raw = list(coll.aggregate(pipeline))
        dept_summary = []
        overall_total_leaves = 0.0
        overall_total_el = 0.0
        overall_total_nh = 0.0
        overall_total_sl = 0.0
        overall_total_eligible = 0.0
        overall_total_taken = 0.0
        overall_total_emps = 0

        for d in dept_summary_raw:
            d_dept = d['_id'] or 'Uncategorized'
            c = d['count']
            tot = round(d.get('total_leaves', 0.0), 2)
            el = round(d.get('total_el', 0.0), 2)
            nh = round(d.get('total_nh', 0.0), 2)
            sl = round(d.get('total_sl', 0.0), 2)
            elig = round(d.get('total_eligible', 0.0), 2)
            taken = round(d.get('total_taken', 0.0), 2)

            dept_summary.append({
                'department': d_dept,
                'department_code': d.get('department_code', ''),
                'employee_count': c,
                'total_leaves': tot,
                'total_el': el,
                'total_nh': nh,
                'total_sl': sl,
                'total_eligible': elig,
                'total_taken': taken,
                'avg_leaves_per_emp': round(tot / c, 1) if c else 0.0
            })

            overall_total_leaves += tot
            overall_total_el += el
            overall_total_nh += nh
            overall_total_sl += sl
            overall_total_eligible += elig
            overall_total_taken += taken
            overall_total_emps += c

        summary = {
            'total_employees': overall_total_emps,
            'total_eligible_leaves': round(overall_total_eligible, 2),
            'total_taken_leaves': round(overall_total_taken, 2),
            'total_available_leaves': round(overall_total_leaves, 2),
            'total_el': round(overall_total_el, 2),
            'total_nh': round(overall_total_nh, 2),
            'total_sl': round(overall_total_sl, 2),
            'filtered_count': len(records),
            'departments': dept_summary
        }

        return Response({
            'isHOD': is_hod,
            'assignedDepartments': hod_assigned_depts if is_hod else [],
            'summary': summary,
            'records': records
        }, status=status.HTTP_200_OK)

    except Exception as e:
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


def resolve_employee_leave_quota(profile_doc, db=None):
    """
    Evaluates active policy rules from 'leave_policy_rules' to determine the best matching quota
    for an employee profile based on department, designation, and isDoctor flag.
    Returns: (dict(EL, NH, SL, total_eligible), matching_policy_doc)
    """
    import os
    from pymongo import MongoClient

    if db is None:
        from employees.views.common.utils import get_mongo_client
        client = get_mongo_client()
        if not client:
            mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
            client = MongoClient(mongo_uri)
        db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]

    policies_col = db['leave_policy_rules']
    active_policies = list(policies_col.find({'is_active': True}))

    default_quota = {'EL': 12.0, 'NH': 9.0, 'SL': 12.0, 'total_eligible': 33.0}
    if not active_policies:
        return default_quota, None

    emp_dept = str(profile_doc.get('department') or '').strip().lower()
    emp_desig = str(profile_doc.get('designation') or '').strip().lower()
    is_doctor = bool(profile_doc.get('isDoctor'))

    scored_policies = []
    for pol in active_policies:
        pol_dept = str(pol.get('department_code') or 'ALL').strip().lower()
        pol_desig = str(pol.get('designation_code') or 'ALL').strip().lower()
        doc_app = str(pol.get('doctor_applicability') or 'ALL').upper().strip()

        # Doctor match check
        if doc_app == 'DOCTOR_ONLY' and not is_doctor:
            continue
        if doc_app == 'NON_DOCTOR_ONLY' and is_doctor:
            continue

        # Department match check
        if pol_dept != 'all' and pol_dept != emp_dept:
            continue

        # Designation match check
        if pol_desig != 'all' and pol_desig != emp_desig:
            continue

        # Specificity score
        score = 0
        if pol_dept != 'all':
            score += 100
        if pol_desig != 'all':
            score += 50
        if doc_app != 'ALL':
            score += 25
        if pol.get('is_default'):
            score += 1

        scored_policies.append((score, pol))

    if not scored_policies:
        for pol in active_policies:
            if pol.get('is_default'):
                el = float(pol.get('el_quota', 12.0))
                nh = float(pol.get('nh_quota', 9.0))
                sl = float(pol.get('sl_quota', 12.0))
                return {'EL': el, 'NH': nh, 'SL': sl, 'total_eligible': round(el + nh + sl, 2)}, pol
        return default_quota, None

    scored_policies.sort(key=lambda x: x[0], reverse=True)
    best_policy = scored_policies[0][1]

    el = float(best_policy.get('el_quota', 12.0))
    nh = float(best_policy.get('nh_quota', 9.0))
    sl = float(best_policy.get('sl_quota', 12.0))

    # Pro-Rata allocation based on 'EL and NH details to IT main.xlsx' for mid-year joiners
    doj_val = profile_doc.get('doj')
    if doj_val and not is_doctor:
        try:
            doj_dt = None
            if isinstance(doj_val, datetime):
                doj_dt = doj_val
            elif isinstance(doj_val, str):
                for fmt in ['%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S']:
                    try:
                        doj_dt = datetime.strptime(doj_val.strip(), fmt)
                        break
                    except ValueError:
                        pass
            if doj_dt and doj_dt.year == datetime.now().year:
                m = doj_dt.month
                # Master Excel month-wise prorata quota schedule
                excel_prorata_map = {
                    1: (12.0, 9.0, 12.0),
                    2: (11.0, 7.0, 11.0),
                    3: (10.0, 7.0, 10.0),
                    4: (9.0, 7.0, 9.0),
                    5: (8.0, 5.0, 8.0),
                    6: (7.0, 5.0, 7.0),
                    7: (6.0, 5.0, 6.0),
                    8: (5.0, 5.0, 5.0),
                    9: (4.0, 4.0, 4.0),
                    10: (3.0, 3.0, 3.0),
                    11: (2.0, 2.0, 2.0),
                    12: (1.0, 1.0, 1.0),
                }
                if m in excel_prorata_map:
                    el, nh, sl = excel_prorata_map[m]
        except Exception:
            pass

    return {'EL': el, 'NH': nh, 'SL': sl, 'total_eligible': round(el + nh + sl, 2)}, best_policy


def provision_employee_leave_balance(emp_id, profile_doc=None, force_recalculate=False, client=None):
    """
    Creates or updates the leave balance record in 'employees_leave_balance' for the given employee.
    Calculates eligible entitlement from leave policy rules and subtracts approved leave taken.
    """
    import os
    from datetime import datetime
    from pymongo import MongoClient

    try:
        if client is None:
            from employees.views.common.utils import get_mongo_client
            client = get_mongo_client()
            if not client:
                mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
                client = MongoClient(mongo_uri)

        db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]
        coll = db['employees_leave_balance']

        if not profile_doc:
            profile_doc = db['backend_diagnostics_profile'].find_one({'employeeId': str(emp_id).strip()})

        if not profile_doc:
            return None

        target_id = str(profile_doc.get('employeeId') or emp_id).strip()
        emp_name = profile_doc.get('employeeName') or f"{profile_doc.get('first_name', '')} {profile_doc.get('last_name', '')}".strip() or target_id
        raw_dept = profile_doc.get('department') or 'DEPT001'
        raw_desig = profile_doc.get('designation') or 'Staff'
        is_doc = bool(profile_doc.get('isDoctor'))

        # Resolve department friendly name
        dept_obj = db['backend_diagnostics_Departments'].find_one({'$or': [{'department_code': raw_dept}, {'department_name': raw_dept}]})
        dept_name = dept_obj.get('department_name') if dept_obj else raw_dept

        # Resolve designation friendly name
        desig_obj = db['backend_diagnostics_Designation'].find_one({'$or': [{'Designation_code': raw_desig}, {'designation': raw_desig}]})
        desig_name = desig_obj.get('designation') if desig_obj else raw_desig

        quota, matched_policy = resolve_employee_leave_quota(profile_doc, db=db)

        # Approved leaves taken calculation
        taken = {'EL': 0.0, 'NH': 0.0, 'SL': 0.0, 'total_taken': 0.0}
        try:
            from employees.models import LeaveRequest
            approved_leaves = LeaveRequest.objects.filter(employee_id=target_id, status='Approved')
            for req in approved_leaves:
                days = (req.end_date - req.start_date).days + 1 if req.start_date and req.end_date else 1
                ltype = str(req.leave_type or '').upper().strip()
                tkey = 'EL'
                if 'NH' in ltype or 'HOLIDAY' in ltype or 'PH' in ltype:
                    tkey = 'NH'
                elif 'SL' in ltype or 'SICK' in ltype or 'CL' in ltype or 'CASUAL' in ltype:
                    tkey = 'SL'
                taken['total_taken'] += days
                taken[tkey] += days
        except Exception as e:
            print(f"Error calculating taken leaves for {target_id}: {e}")

        el_bal = max(0.0, quota['EL'] - taken['EL'])
        nh_bal = max(0.0, quota['NH'] - taken['NH'])
        sl_bal = max(0.0, quota['SL'] - taken['SL'])
        total_avail = round(el_bal + nh_bal + sl_bal, 2)

        now = datetime.now()
        existing = coll.find_one({'employee_id': target_id})

        if not existing:
            doc_to_save = {
                'employee_id': target_id,
                'employee_code': target_id,
                'employee_name': emp_name,
                'department': dept_name,
                'department_code': raw_dept,
                'designation': desig_name,
                'designation_code': raw_desig,
                'is_doctor': is_doc,
                'doj': profile_doc.get('doj') or '2024-01-01',
                'as_of_year': now.year,
                'policy_id': str(matched_policy.get('_id')) if matched_policy else None,
                'policy_name': matched_policy.get('policy_name') if matched_policy else 'Default Policy',
                'eligible_entitlement': quota,
                'leave_balances': {
                    'EL': el_bal,
                    'NH': nh_bal,
                    'SL': sl_bal,
                    'total_available': total_avail
                },
                'leaves_taken': taken,
                'created_at': now,
                'updated_at': now
            }
            coll.insert_one(doc_to_save)
            return doc_to_save
        else:
            update_payload = {
                'employee_name': emp_name,
                'department': dept_name,
                'department_code': raw_dept,
                'designation': desig_name,
                'designation_code': raw_desig,
                'is_doctor': is_doc,
                'updated_at': now
            }
            if force_recalculate:
                update_payload['policy_id'] = str(matched_policy.get('_id')) if matched_policy else None
                update_payload['policy_name'] = matched_policy.get('policy_name') if matched_policy else 'Default Policy'
                update_payload['eligible_entitlement'] = quota
                update_payload['leave_balances'] = {
                    'EL': el_bal,
                    'NH': nh_bal,
                    'SL': sl_bal,
                    'total_available': total_avail
                }
                update_payload['leaves_taken'] = taken

            coll.update_one({'_id': existing['_id']}, {'$set': update_payload})
            return coll.find_one({'_id': existing['_id']})
    except Exception as exc:
        print(f"Error provisioning leave balance for {emp_id}: {exc}")
        return None


@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
def leave_policies(request):
    """
    GET: List all configured leave policy rules, along with department/designation lookup lists and stats.
    POST: Create a new leave policy rule.
    """
    import os
    from bson import ObjectId
    from pymongo import MongoClient

    mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
    client = MongoClient(mongo_uri)
    db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]
    policies_col = db['leave_policy_rules']

    if request.method == 'GET':
        policies = list(policies_col.find().sort([('is_default', -1), ('created_at', -1)]))
        for p in policies:
            p['id'] = str(p['_id'])
            p['_id'] = str(p['_id'])
            if 'created_at' in p and hasattr(p['created_at'], 'isoformat'):
                p['created_at'] = p['created_at'].isoformat()
            if 'updated_at' in p and hasattr(p['updated_at'], 'isoformat'):
                p['updated_at'] = p['updated_at'].isoformat()

        # Lookup options for Department and Designation
        depts_raw = list(db['backend_diagnostics_Departments'].find({}, {'_id': 0, 'department_code': 1, 'department_name': 1}))
        dept_options = [{'code': d.get('department_code'), 'name': d.get('department_name') or d.get('department_code')} for d in depts_raw if d.get('department_code')]
        dept_options.sort(key=lambda x: x['name'])

        desigs_raw = list(db['backend_diagnostics_Designation'].find({}, {'_id': 0, 'Designation_code': 1, 'designation': 1}))
        desig_options = [{'code': d.get('Designation_code'), 'name': d.get('designation') or d.get('Designation_code')} for d in desigs_raw if d.get('Designation_code')]
        desig_options.sort(key=lambda x: x['name'])

        total_policies = len(policies)
        active_policies = sum(1 for p in policies if p.get('is_active', True))
        total_synced_employees = db['employees_leave_balance'].count_documents({})
        doctor_policies = sum(1 for p in policies if p.get('doctor_applicability') == 'DOCTOR_ONLY')

        prorata_schedule = [
            {'month': 1, 'month_name': 'January', 'el': 12.0, 'nh': 9.0, 'sl': 12.0, 'total': 33.0},
            {'month': 2, 'month_name': 'February', 'el': 11.0, 'nh': 7.0, 'sl': 11.0, 'total': 29.0},
            {'month': 3, 'month_name': 'March', 'el': 10.0, 'nh': 7.0, 'sl': 10.0, 'total': 27.0},
            {'month': 4, 'month_name': 'April', 'el': 9.0, 'nh': 7.0, 'sl': 9.0, 'total': 25.0},
            {'month': 5, 'month_name': 'May', 'el': 8.0, 'nh': 5.0, 'sl': 8.0, 'total': 21.0},
            {'month': 6, 'month_name': 'June', 'el': 7.0, 'nh': 5.0, 'sl': 7.0, 'total': 19.0},
            {'month': 7, 'month_name': 'July', 'el': 6.0, 'nh': 5.0, 'sl': 6.0, 'total': 17.0},
            {'month': 8, 'month_name': 'August', 'el': 5.0, 'nh': 5.0, 'sl': 5.0, 'total': 15.0},
            {'month': 9, 'month_name': 'September', 'el': 4.0, 'nh': 4.0, 'sl': 4.0, 'total': 12.0},
            {'month': 10, 'month_name': 'October', 'el': 3.0, 'nh': 3.0, 'sl': 3.0, 'total': 9.0},
            {'month': 11, 'month_name': 'November', 'el': 2.0, 'nh': 2.0, 'sl': 2.0, 'total': 6.0},
            {'month': 12, 'month_name': 'December', 'el': 1.0, 'nh': 1.0, 'sl': 1.0, 'total': 3.0},
        ]

        excel_path = '/Users/parthibanmurugan/Desktop/Live Projects/HR/EL and NH details to IT main.xlsx'
        excel_info = {
            'exists': os.path.exists(excel_path),
            'filename': 'EL and NH details to IT main.xlsx',
            'total_employees': 274,
            'total_eligible_leaves': 7763.0,
            'total_used_leaves': 1351.0,
            'total_pending_leaves': 6412.0
        }

        return Response({
            'policies': policies,
            'departments': dept_options,
            'designations': desig_options,
            'prorata_schedule': prorata_schedule,
            'excel_info': excel_info,
            'stats': {
                'total_policies': total_policies,
                'active_policies': active_policies,
                'total_synced_employees': total_synced_employees,
                'doctor_policies': doctor_policies
            }
        }, status=status.HTTP_200_OK)

    elif request.method == 'POST':
        data = request.data
        policy_name = str(data.get('policy_name', '')).strip()
        if not policy_name:
            return Response({'error': 'Policy Name is required.'}, status=status.HTTP_400_BAD_REQUEST)

        dept_code = str(data.get('department_code') or 'ALL').strip()
        dept_name = str(data.get('department_name') or 'All Departments').strip()
        desig_code = str(data.get('designation_code') or 'ALL').strip()
        desig_name = str(data.get('designation_name') or 'All Designations').strip()
        doctor_app = str(data.get('doctor_applicability') or 'ALL').strip().upper()
        if doctor_app not in ['ALL', 'DOCTOR_ONLY', 'NON_DOCTOR_ONLY']:
            doctor_app = 'ALL'

        try:
            el_quota = float(data.get('el_quota', 12.0))
            nh_quota = float(data.get('nh_quota', 9.0))
            sl_quota = float(data.get('sl_quota', 12.0))
        except (ValueError, TypeError):
            return Response({'error': 'Quota values must be valid numbers.'}, status=status.HTTP_400_BAD_REQUEST)

        total_quota = round(el_quota + nh_quota + sl_quota, 2)
        is_active = bool(data.get('is_active', True))
        is_default = bool(data.get('is_default', False))
        description = str(data.get('description', '')).strip()

        now = datetime.now()
        new_policy = {
            'policy_name': policy_name,
            'department_code': dept_code,
            'department_name': dept_name,
            'designation_code': desig_code,
            'designation_name': desig_name,
            'doctor_applicability': doctor_app,
            'el_quota': el_quota,
            'nh_quota': nh_quota,
            'sl_quota': sl_quota,
            'total_quota': total_quota,
            'description': description,
            'is_active': is_active,
            'is_default': is_default,
            'created_at': now,
            'updated_at': now,
            'created_by': request.data.get('auth_user', 'admin')
        }

        result = policies_col.insert_one(new_policy)
        new_policy['id'] = str(result.inserted_id)
        new_policy['_id'] = str(result.inserted_id)
        new_policy['created_at'] = now.isoformat()
        new_policy['updated_at'] = now.isoformat()

        return Response({
            'success': True,
            'message': f"Policy '{policy_name}' created successfully.",
            'policy': new_policy
        }, status=status.HTTP_201_CREATED)


@api_view(['PUT', 'DELETE'])
@permission_classes([AllowAny])
def leave_policy_detail(request, policy_id):
    """
    PUT: Update an existing policy rule.
    DELETE: Delete a policy rule (unless it is the protected system default).
    """
    import os
    from bson import ObjectId
    from pymongo import MongoClient

    mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
    client = MongoClient(mongo_uri)
    db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]
    policies_col = db['leave_policy_rules']

    try:
        oid = ObjectId(policy_id)
    except Exception:
        return Response({'error': 'Invalid Policy ID format.'}, status=status.HTTP_400_BAD_REQUEST)

    existing = policies_col.find_one({'_id': oid})
    if not existing:
        return Response({'error': 'Policy rule not found.'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'DELETE':
        if existing.get('is_default'):
            return Response({'error': 'Cannot delete the system Global Standard Default Policy.'}, status=status.HTTP_400_BAD_REQUEST)
        policies_col.delete_one({'_id': oid})
        return Response({'success': True, 'message': 'Policy rule deleted successfully.'}, status=status.HTTP_200_OK)

    elif request.method == 'PUT':
        data = request.data
        update_fields = {'updated_at': datetime.now()}

        if 'policy_name' in data:
            update_fields['policy_name'] = str(data['policy_name']).strip()
        if 'department_code' in data:
            update_fields['department_code'] = str(data['department_code']).strip()
        if 'department_name' in data:
            update_fields['department_name'] = str(data['department_name']).strip()
        if 'designation_code' in data:
            update_fields['designation_code'] = str(data['designation_code']).strip()
        if 'designation_name' in data:
            update_fields['designation_name'] = str(data['designation_name']).strip()
        if 'doctor_applicability' in data:
            doc_app = str(data['doctor_applicability']).strip().upper()
            if doc_app in ['ALL', 'DOCTOR_ONLY', 'NON_DOCTOR_ONLY']:
                update_fields['doctor_applicability'] = doc_app
        if 'description' in data:
            update_fields['description'] = str(data['description']).strip()
        if 'is_active' in data:
            update_fields['is_active'] = bool(data['is_active'])

        el_val = float(data.get('el_quota', existing.get('el_quota', 12.0)))
        nh_val = float(data.get('nh_quota', existing.get('nh_quota', 9.0)))
        sl_val = float(data.get('sl_quota', existing.get('sl_quota', 12.0)))
        update_fields['el_quota'] = el_val
        update_fields['nh_quota'] = nh_val
        update_fields['sl_quota'] = sl_val
        update_fields['total_quota'] = round(el_val + nh_val + sl_val, 2)

        policies_col.update_one({'_id': oid}, {'$set': update_fields})
        updated_doc = policies_col.find_one({'_id': oid})
        updated_doc['id'] = str(updated_doc['_id'])
        updated_doc['_id'] = str(updated_doc['_id'])
        if 'created_at' in updated_doc and hasattr(updated_doc['created_at'], 'isoformat'):
            updated_doc['created_at'] = updated_doc['created_at'].isoformat()
        if 'updated_at' in updated_doc and hasattr(updated_doc['updated_at'], 'isoformat'):
            updated_doc['updated_at'] = updated_doc['updated_at'].isoformat()

        return Response({
            'success': True,
            'message': 'Policy rule updated successfully.',
            'policy': updated_doc
        }, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([AllowAny])
def sync_all_leave_balances(request):
    """
    POST: Synchronizes/re-provisions leave balances for all active employee profiles
    based on master Excel 'EL and NH details to IT main.xlsx' and active policy rules.
    """
    import os
    from datetime import datetime
    from pymongo import MongoClient

    mongo_uri = os.environ.get('GLOBAL_DB_HOST', 'mongodb://admin:SMRFT%40prod2026@45.252.190.162:27017/')
    client = MongoClient(mongo_uri)
    db = client[os.environ.get('GLOBAL_DB_NAME', 'Global')]
    coll = db['employees_leave_balance']

    department_filter = request.data.get('department_code')
    excel_synced_count = 0
    excel_emp_codes = set()

    # 1. Sync from master Excel file 'EL and NH details to IT main.xlsx' if exists
    excel_path = '/Users/parthibanmurugan/Desktop/Live Projects/HR/EL and NH details to IT main.xlsx'
    if os.path.exists(excel_path) and (not department_filter or department_filter == 'ALL'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb['Final']
            now = datetime.now()

            for r in range(3, ws.max_row + 1):
                val = ws.cell(r, 2).value
                if not val:
                    continue
                code = str(val).strip()
                excel_emp_codes.add(code)
                name = ws.cell(r, 4).value
                desig = ws.cell(r, 6).value
                doj_val = ws.cell(r, 3).value

                used_nh = float(ws.cell(r, 7).value or 0)
                used_el = float(ws.cell(r, 8).value or 0)
                used_sl = float(ws.cell(r, 9).value or 0)

                elig_nh = float(ws.cell(r, 10).value or 0)
                elig_el = float(ws.cell(r, 11).value or 0)
                elig_sl = float(ws.cell(r, 12).value or 0)

                pend_nh = float(ws.cell(r, 13).value or 0)
                pend_el = float(ws.cell(r, 14).value or 0)
                pend_sl = float(ws.cell(r, 15).value or 0)

                existing = coll.find_one({'employee_id': code})
                payload = {
                    'employee_id': code,
                    'employee_code': code,
                    'employee_name': name or (existing.get('employee_name') if existing else code),
                    'designation': desig or (existing.get('designation') if existing else ''),
                    'as_of_year': now.year,
                    'source_file': 'EL and NH details to IT main.xlsx',
                    'eligible_entitlement': {
                        'EL': elig_el,
                        'NH': elig_nh,
                        'SL': elig_sl,
                        'total_eligible': round(elig_el + elig_nh + elig_sl, 2)
                    },
                    'leaves_taken': {
                        'EL': used_el,
                        'NH': used_nh,
                        'SL': used_sl,
                        'total_taken': round(used_el + used_nh + used_sl, 2)
                    },
                    'leave_balances': {
                        'EL': pend_el,
                        'NH': pend_nh,
                        'SL': pend_sl,
                        'total_available': round(pend_el + pend_nh + pend_sl, 2)
                    },
                    'updated_at': now
                }
                if doj_val:
                    payload['doj'] = doj_val.strftime('%Y-%m-%d') if isinstance(doj_val, datetime) else str(doj_val)

                if existing:
                    coll.update_one({'_id': existing['_id']}, {'$set': payload})
                else:
                    payload['created_at'] = now
                    coll.insert_one(payload)
                excel_synced_count += 1
        except Exception as ex_err:
            print("Error syncing master excel in sync_all_leave_balances:", ex_err)

    # 2. Sync remaining profiles in backend_diagnostics_profile (Doctors and new joiners)
    query = {}
    if department_filter and department_filter != 'ALL':
        query['department'] = department_filter

    profiles = list(db['backend_diagnostics_profile'].find(query))
    synced_count = 0
    errors = []

    for prof in profiles:
        emp_id = str(prof.get('employeeId') or '').strip()
        if not emp_id or emp_id in excel_emp_codes:
            continue
        try:
            provision_employee_leave_balance(emp_id, profile_doc=prof, force_recalculate=True, client=client)
            synced_count += 1
        except Exception as e:
            errors.append(f"{emp_id}: {str(e)}")

    total_synced = excel_synced_count + synced_count
    return Response({
        'success': True,
        'message': f"Successfully synchronized {total_synced} employee leave balances ({excel_synced_count} from master Excel + {synced_count} profiles/doctors from policy rules).",
        'excel_synced_count': excel_synced_count,
        'profiles_synced_count': synced_count,
        'total_synced': total_synced,
        'errors_count': len(errors)
    }, status=status.HTTP_200_OK)

